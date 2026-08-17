"""One-time, reproducible extraction of CREST's real active-occupant-count
transition-probability matrices (TPMs) into
``src/occupancy/households/data/tpm_crest.json``.

Not part of the installed ``occupancy`` package -- a dev-only tool, run
once (or re-run if the source workbook is ever replaced), mirroring how
``households/data/equipment.json``'s CREST-sourced fields were originally
extracted "directly by the user," not via a tracked pipeline dependency.

**Requires** ``openpyxl`` (not an ``occupancy`` runtime/dev dependency --
install ad hoc: ``pip install openpyxl`` or ``conda install -c conda-forge
openpyxl`` into the dev env before running this script).

**Source**: ``data/inputs/CREST_Domestic_electricity_demand_model_1.0e.xlsm``
(Richardson, Thomson, Infield -- Loughborough University CREST), gitignored
and not committed to this repo (CC BY-NC-ND -- no redistribution of the
raw workbook). Obtain it directly via the paper's own public download link
if it's missing locally; see ``.claude/residential/resolved.md`` for the
provenance/licensing notes already established for this repo's other
CREST-sourced data (``households/data/equipment.json``).

**What this reads**: 10 sheets, ``tpm{1..5}_{wd,we}`` -- the "1-5" indexes
household size (1-5 residents, confirmed via the workbook's own "main"
sheet resident-count picker and each tpm sheet's "One/Two/.../Five
resident house" row-5 label), each holding a 144-period (10-minute
resolution) x 7x7 (active-occupant states 0-6) transition matrix in rows
11-1018 / columns C-I.

Usage::

    python scripts/extract_crest_tpm.py
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import numpy as np

# Imported from the installed package (this dev env has `occupancy`
# installed editable, per infrastructure/env/occupancy_env.yml's `pip
# install -e . --no-deps` step) so the composition math isn't duplicated
# between this script and households/crest_tpm.py.
from occupancy.households.crest_tpm import compose_hourly_transition_matrix

_REPO_ROOT = Path(__file__).resolve().parent.parent
_XLSM_PATH = (
    _REPO_ROOT
    / "data"
    / "inputs"
    / "CREST_Domestic_electricity_demand_model_1.0e.xlsm"
)
_OUTPUT_PATH = (
    _REPO_ROOT / "src" / "occupancy" / "households" / "data" / "tpm_crest.json"
)

_N_STATES_RAW = 7  # active-occupant states 0-6, as stored in the workbook
_N_PERIODS = 144  # 10-minute periods/day
_DATA_FIRST_ROW = 11  # 1-indexed openpyxl row of period=1, from-state=0
_DATA_FIRST_COL = 3  # column C: probability of transition to state 0

# Regression anchor: the real value at tpm1_wd, period 1, from-state 0 (the
# first data row of the first sheet), confirmed by direct inspection this
# session. Catches a row/column-offset error in the read before it can
# silently corrupt the extracted data.
_ANCHOR_SHEET = "tpm1_wd"
_ANCHOR_ROW_PROBS = np.array(
    [0.99442896935933101, 5.5710306406685202e-3, 0.0, 0.0, 0.0, 0.0, 0.0]
)


def _read_sheet(ws: Any) -> np.ndarray:
    """Read one tpm{m}_{wd,we} sheet into (144, 7, 7): period x from-state
    x to-state.

    Uses ``iter_rows(..., values_only=True)`` -- openpyxl's read-only mode
    is optimized for sequential row iteration, not random ``ws.cell(row=,
    column=)`` access (which re-scans from the start of the underlying XML
    stream on every call and is pathologically slow for ~1000-row sheets
    accessed cell-by-cell).
    """
    last_row = _DATA_FIRST_ROW + _N_PERIODS * _N_STATES_RAW - 1
    last_col = _DATA_FIRST_COL + _N_STATES_RAW - 1
    matrix = np.zeros((_N_PERIODS, _N_STATES_RAW, _N_STATES_RAW))
    rows = ws.iter_rows(
        min_row=_DATA_FIRST_ROW,
        max_row=last_row,
        min_col=_DATA_FIRST_COL,
        max_col=last_col,
        values_only=True,
    )
    for flat_index, row_values in enumerate(rows):
        period, from_state = divmod(flat_index, _N_STATES_RAW)
        matrix[period, from_state, :] = [
            float(v) if v is not None else 0.0 for v in row_values
        ]
    return matrix


def _validate_and_trim(
    matrix: np.ndarray, num_persons: int, sheet_name: str
) -> tuple[np.ndarray, int]:
    """Sanity-check raw (144, 7, 7) data, then trim to the reachable
    (num_persons + 1, num_persons + 1) submatrix. Returns (trimmed,
    unobserved_row_count).

    Only *reachable* from-states (0..num_persons -- a household of this
    size can never actually have more active occupants than residents) are
    considered. Unreachable from-states (num_persons+1..6) are expected to
    be entirely blank/zero in the source workbook -- checked, not assumed
    -- since those rows describe a state that can never occur for this
    household size.

    Among reachable rows, a genuinely *unobserved* state at a given
    10-minute period (e.g. "all N residents simultaneously active at
    01:00") legitimately sums to exactly 0 in CREST's own data -- not an
    extraction bug, just no survey observations to estimate a transition
    from. These are replaced with an identity/self-loop row (stay in the
    same state with probability 1): a defensible, clearly-flagged fallback
    for an unobserved state, not a silently invented one -- if the state
    is never observed, nothing in the data suggests occupants would leave
    it, and self-loop is the natural "no information" default (the
    returned count lets the caller confirm how many such rows existed
    rather than this passing silently). A reachable row summing to
    anything else (neither ~1 nor ~0) is a genuine anomaly and still
    raises.
    """
    n = num_persons + 1

    reachable = matrix[:, :n, :].copy()
    reachable_row_sums = reachable.sum(axis=-1)
    is_unit = np.isclose(reachable_row_sums, 1.0, atol=1e-6)
    is_zero = np.isclose(reachable_row_sums, 0.0, atol=1e-9)
    bad = np.argwhere(~is_unit & ~is_zero)
    if len(bad) > 0:
        p, s = bad[0]
        raise ValueError(
            f"{sheet_name}: {len(bad)} reachable row(s) sum to neither 0 "
            f"nor 1 (e.g. period={p}, from_state={s}, "
            f"sum={reachable_row_sums[p, s]!r})"
        )

    unobserved_row_count = int((~is_unit & is_zero).sum())
    for p, s in np.argwhere(~is_unit & is_zero):
        reachable[p, s, :] = 0.0
        reachable[p, s, s] = 1.0  # self-loop fallback

    unreachable_out = matrix[:, :n, n:]
    unreachable_from_states = matrix[:, n:, :]
    if not np.allclose(unreachable_out, 0.0, atol=1e-3):
        raise ValueError(
            f"{sheet_name}: unreachable to-states (>{num_persons} active "
            f"occupants in a {num_persons}-resident house) carry "
            f"non-negligible probability mass (max="
            f"{unreachable_out.max()!r}) -- household-size assumption "
            "may be wrong."
        )
    if not np.allclose(unreachable_from_states, 0.0, atol=1e-6):
        raise ValueError(
            f"{sheet_name}: unreachable from-states (>{num_persons} "
            f"active occupants) carry non-negligible probability mass "
            f"(max={unreachable_from_states.max()!r}) -- expected these "
            "to be entirely blank/zero in the source workbook."
        )

    trimmed = reachable[:, :, :n]
    # Re-normalize to sum to exactly 1 (dropped to-state columns may have
    # carried tiny residual probability mass within the 1e-3 tolerance
    # above; self-loop rows already sum to exactly 1 and are unaffected).
    trimmed = trimmed / trimmed.sum(axis=-1, keepdims=True)
    return trimmed, unobserved_row_count


def main() -> None:
    try:
        import openpyxl
    except ImportError as exc:
        raise SystemExit(
            "openpyxl is required to run this script -- install ad hoc "
            "(`pip install openpyxl` or `conda install -c conda-forge "
            "openpyxl`), it is not an occupancy runtime/dev dependency."
        ) from exc

    if not _XLSM_PATH.exists():
        raise SystemExit(
            f"CREST workbook not found at {_XLSM_PATH}. Obtain it "
            "directly via the CREST Domestic Electricity Demand Model "
            "1.0e's own public download link (Richardson, Thomson, "
            "Infield -- Loughborough University CREST); see "
            ".claude/residential/resolved.md for provenance notes."
        )

    print(f"Reading {_XLSM_PATH} ...")
    workbook = openpyxl.load_workbook(
        _XLSM_PATH, read_only=True, data_only=True
    )

    # Regression anchor check before trusting anything else in the file.
    anchor_ws = workbook[_ANCHOR_SHEET]
    anchor_row = [
        anchor_ws.cell(row=_DATA_FIRST_ROW, column=_DATA_FIRST_COL + c).value
        for c in range(_N_STATES_RAW)
    ]
    anchor_row_arr = np.array(
        [float(v) if v is not None else 0.0 for v in anchor_row]
    )
    if not np.allclose(anchor_row_arr, _ANCHOR_ROW_PROBS, atol=1e-9):
        raise SystemExit(
            f"Regression anchor check failed: {_ANCHOR_SHEET} period 1 "
            f"from-state 0 read {anchor_row_arr.tolist()}, expected "
            f"{_ANCHOR_ROW_PROBS.tolist()}. The row/column offsets "
            "(_DATA_FIRST_ROW/_DATA_FIRST_COL) may no longer match this "
            "workbook's layout -- aborting rather than emitting silently "
            "wrong data."
        )
    print("Regression anchor check passed.")

    output: dict[str, Any] = {
        "_comment": (
            "Hourly-composed CREST active-occupant-count transition-"
            "probability matrices, per household size (1-5 residents) "
            "and day-type. Composed from CREST's native 10-minute (144 "
            "steps/day) tpm<m>_wd/we sheets by matrix-multiplying the 6 "
            "consecutive 10-minute matrices covering each hour (see "
            "occupancy.households.crest_tpm."
            "compose_hourly_transition_matrix). Not applicable above 5 "
            "residents -- see households/crest_tpm.py's "
            "MAX_CALIBRATED_SIZE fallback (households/household_profile.py "
            "falls back to the synthesized markov_chain generator for "
            "larger households). A handful of reachable (from_state <= "
            "num_persons) 10-minute rows have zero survey observations "
            "in CREST's own data (e.g. 'all N residents simultaneously "
            "active at 01:00' for larger N) and are filled with a "
            "self-loop (stay in the same state with probability 1) "
            "rather than left undefined -- see "
            "scripts/extract_crest_tpm.py's _validate_and_trim "
            "docstring. Generated by scripts/extract_crest_tpm.py "
            "-- do not hand-edit."
        ),
        "_crest_source": (
            "CREST Domestic Electricity Demand Model 1.0e (Richardson, "
            "Thomson, Infield -- Loughborough University CREST), obtained "
            "via the paper's own public download link. Underlying survey: "
            "Ipsos-RSL and Office for National Statistics, United Kingdom "
            "Time Use Survey, 2000 (Computer File), third ed., UK Data "
            "Archive (distributor), Colchester, Essex, September 2003, "
            "SN: 4504 (credited on the workbook's own 'main' sheet)."
        ),
    }

    for num_persons in range(1, 6):
        last_row = _DATA_FIRST_ROW + _N_PERIODS * _N_STATES_RAW - 1
        entry: dict[str, Any] = {
            "_crest_source": (
                f"tpm{num_persons}_wd / tpm{num_persons}_we sheets, rows "
                f"{_DATA_FIRST_ROW}-{last_row} (144 x 7x7), trimmed to "
                f"the reachable {num_persons + 1}x{num_persons + 1} "
                "submatrix after validating the remaining states carry "
                "~0 probability mass, then hourly-composed."
            )
        }
        for day_type, suffix in (("weekday", "wd"), ("weekend", "we")):
            sheet_name = f"tpm{num_persons}_{suffix}"
            ws = workbook[sheet_name]
            raw = _read_sheet(ws)
            trimmed, unobserved = _validate_and_trim(
                raw, num_persons, sheet_name
            )
            print(
                f"  Extracted {sheet_name} "
                f"({unobserved} unobserved reachable state(s) "
                "self-loop-filled)"
            )
            hourly = compose_hourly_transition_matrix(trimmed)
            hourly_row_sums = hourly.sum(axis=-1)
            if not np.allclose(hourly_row_sums, 1.0, atol=1e-6):
                raise ValueError(
                    f"{sheet_name}: composed hourly rows do not sum to 1 "
                    f"(max deviation="
                    f"{np.abs(hourly_row_sums - 1.0).max()!r})"
                )
            entry[day_type] = hourly.tolist()
        output[str(num_persons)] = entry

    workbook.close()

    _OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with _OUTPUT_PATH.open("w", encoding="utf-8") as handle:
        json.dump(output, handle, indent=2)
    print(f"Wrote {_OUTPUT_PATH}")


if __name__ == "__main__":
    main()
